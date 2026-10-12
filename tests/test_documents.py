import io
import zipfile
from pathlib import Path

import docx
import openpyxl
import pytest

from hyperscale.pii_scanner import HitKind
from hyperscale.pii_scanner import scan
from hyperscale.pii_scanner.documents import extract_text
from hyperscale.pii_scanner.documents import ExtractedText
from hyperscale.pii_scanner.documents import Segment


def _docx(
    paragraphs: list[str],
    table: list[list[str]] | None = None,
    header: str | None = None,
) -> bytes:
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    if header is not None:
        document.sections[0].header.paragraphs[0].text = header
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _xlsx(cells: dict[str, dict[str, str | int]]) -> bytes:
    wb = openpyxl.Workbook()
    first = True
    for sheet, values in cells.items():
        ws = wb.active if first else wb.create_sheet()
        assert ws is not None
        ws.title = sheet
        first = False
        for ref, value in values.items():
            ws[ref] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_docx_paragraphs_tables_and_headers_get_labels():
    data = _docx(
        ["Intro.", "", "Jane Smith attended."],
        table=[["Name", "Town"], ["Bob Jones", "Leeds"]],
        header="Case ZC24C50123",
    )
    doc = extract_text(data, kind="docx")
    assert doc.locate(doc.text.index("Jane Smith")) == "paragraph 3"
    assert doc.locate(doc.text.index("Bob Jones")) == "table 1, row 2, cell 1"
    assert doc.locate(doc.text.index("Leeds")) == "table 1, row 2, cell 2"
    assert doc.locate(doc.text.index("ZC24C50123")) == "header 1"
    assert "\n\n" not in doc.text  # empty paragraphs are skipped


def test_xlsx_cells_get_sheet_and_coordinate_labels():
    data = _xlsx(
        {
            "Results": {"A1": "Patient", "B20": "Jane Smith", "C3": 42},
            "Notes": {"A1": "Harrogate"},
        }
    )
    doc = extract_text(data, kind="xlsx")
    assert doc.locate(doc.text.index("Jane Smith")) == "sheet Results, cell B20"
    assert doc.locate(doc.text.index("42")) == "sheet Results, cell C3"
    assert doc.locate(doc.text.index("Harrogate")) == "sheet Notes, cell A1"


def test_hits_map_back_to_locations():
    doc = extract_text(
        _docx(["Nothing here.", "NHS number 943 476 5919."]), kind="docx"
    )
    hits = scan(doc.text)
    nhs = next(h for h in hits if h.kind == HitKind.NHS_NUMBER)
    assert doc.locate(nhs.start) == "paragraph 2"


def test_extract_handles_empty_documents():
    assert extract_text(_docx([]), kind="docx") == ExtractedText(text="", segments=())
    header_only = extract_text(_docx([], header="only header"), kind="docx")
    assert header_only.segments == (Segment(0, 11, "header 1"),)
    assert extract_text(_xlsx({"Empty": {}}), kind="xlsx") == ExtractedText(
        text="", segments=()
    )


def test_locate_separator_offset_maps_to_preceding_segment_and_rejects_out_of_range():
    doc = extract_text(_docx(["One", "Two"]), kind="docx")
    assert doc.text == "One\nTwo"
    assert doc.locate(3) == "paragraph 1"
    with pytest.raises(ValueError):
        doc.locate(99)
    with pytest.raises(ValueError):
        doc.locate(-1)


def test_path_input(tmp_path: Path):
    path = tmp_path / "x.docx"
    path.write_bytes(_docx(["Hello"]))
    assert extract_text(path, kind="docx").text == "Hello"
    assert extract_text(str(path), kind="docx").text == "Hello"


def test_rejects_unknown_kind_and_non_zip_and_oversize(monkeypatch: pytest.MonkeyPatch):
    with pytest.raises(ValueError, match="kind"):
        extract_text(b"x", kind="pdf")  # type: ignore[arg-type]
    with pytest.raises(ValueError, match="zip"):
        extract_text(b"not a zip file", kind="docx")
    monkeypatch.setattr("hyperscale.pii_scanner.documents.MAX_BYTES", 10)
    with pytest.raises(ValueError, match="larger"):
        extract_text(_docx(["Hello"]), kind="docx")


def test_rejects_zip_bombs_by_declared_size(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr("hyperscale.pii_scanner.documents.MAX_UNCOMPRESSED_BYTES", 100)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", "0" * 10_000)
    with pytest.raises(ValueError, match="uncompressed"):
        extract_text(buf.getvalue(), kind="docx")


def test_corrupt_document_raises_value_error():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "not a document")
    with pytest.raises(ValueError, match="could not parse"):
        extract_text(buf.getvalue(), kind="docx")
    with pytest.raises(ValueError, match="could not parse"):
        extract_text(buf.getvalue(), kind="xlsx")


def test_openpyxl_uses_defusedxml():
    import openpyxl.xml

    assert openpyxl.xml.DEFUSEDXML is True
