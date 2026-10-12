"""Optional text extraction from Word and Excel files with a location map.

Install with ``hyperscale-pii-scanner[documents]``. Every file is treated as
untrusted input: size caps, a declared-uncompressed-size check before parsing,
entity resolution disabled in python-docx, defusedxml installed for openpyxl.
"""

from __future__ import annotations

import bisect
import io
import os
import zipfile
from dataclasses import dataclass
from typing import cast
from typing import Literal

# openpyxl uses lxml with its own entity-safe parser when lxml is present
# (python-docx installs it) and defusedxml for streaming parses when installed.
try:
    import docx
    import openpyxl
except ImportError as exc:  # pragma: no cover - exercised only without the extra
    raise ImportError(
        "hyperscale.pii_scanner.documents needs the 'documents' extra: "
        "uv add 'hyperscale-pii-scanner[documents]'"
    ) from exc

from docx.document import Document as DocumentObject
from docx.oxml.document import CT_Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.section import Section
from docx.text.paragraph import Paragraph
from docx.types import ProvidesStoryPart

MAX_BYTES = 50 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
SEPARATOR = "\n"

Kind = Literal["docx", "xlsx"]


@dataclass(frozen=True, slots=True)
class Segment:
    start: int
    end: int
    label: str


@dataclass(frozen=True, slots=True)
class ExtractedText:
    text: str
    segments: tuple[Segment, ...]

    def locate(self, offset: int) -> str:
        """Label of the segment containing ``offset``.

        A separator offset maps to the preceding segment.
        """
        if offset < 0 or offset > len(self.text) or not self.segments:
            raise ValueError(f"offset {offset} is outside the extracted text")
        starts = [s.start for s in self.segments]
        index = bisect.bisect_right(starts, offset) - 1
        if index < 0:
            raise ValueError(f"offset {offset} precedes the first segment")
        return self.segments[index].label


def _read_source(source: str | os.PathLike[str] | bytes) -> bytes:
    if isinstance(source, bytes):
        data = source
    else:
        path = os.fspath(source)
        if os.path.getsize(path) > MAX_BYTES:
            raise ValueError(f"file is larger than {MAX_BYTES} bytes")
        with open(path, "rb") as handle:
            data = handle.read()
    if len(data) > MAX_BYTES:
        raise ValueError(f"file is larger than {MAX_BYTES} bytes")
    return data


def _check_zip(data: bytes) -> None:
    buffer = io.BytesIO(data)
    if not zipfile.is_zipfile(buffer):
        raise ValueError("file is not a zip container (expected .docx or .xlsx)")
    with zipfile.ZipFile(buffer) as zf:
        declared = sum(info.file_size for info in zf.infolist())
    if declared > MAX_UNCOMPRESSED_BYTES:
        raise ValueError(
            f"declared uncompressed size {declared} exceeds "
            f"{MAX_UNCOMPRESSED_BYTES} bytes"
        )


class _Builder:
    def __init__(self) -> None:
        self.parts: list[str] = []
        self.segments: list[Segment] = []
        self.length = 0

    def add(self, text: str, label: str) -> None:
        text = text.strip()
        if not text:
            return
        if self.parts:
            self.length += len(SEPARATOR)
        start = self.length
        self.parts.append(text)
        self.length += len(text)
        self.segments.append(Segment(start, self.length, label))

    def build(self) -> ExtractedText:
        return ExtractedText(SEPARATOR.join(self.parts), tuple(self.segments))


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _walk_wordml(
    root: BaseOxmlElement,
    parent: ProvidesStoryPart,
    builder: _Builder,
    *,
    paragraph_label: str | None,
    table_prefix: str,
) -> None:
    """Add every paragraph under ``root`` in document order.

    Walking the XML (rather than python-docx's block iterators) reaches
    paragraphs inside content controls and nested tables. Body paragraphs are
    numbered unless ``paragraph_label`` is given (headers and footers use one
    label). A paragraph inside a table cell is labelled by its innermost table,
    row and cell; merged cells appear once because the walk follows the XML,
    not the grid.
    """
    paragraph_number = 0
    table_numbers: dict[str, int] = {}
    tree = root.getroottree()
    for p in root.iter(_W + "p"):
        cell = row = table = None
        for ancestor in p.iterancestors():
            if ancestor is root:
                break
            if cell is None and ancestor.tag == _W + "tc":
                cell = ancestor
            elif cell is not None and row is None and ancestor.tag == _W + "tr":
                row = ancestor
            elif row is not None and ancestor.tag == _W + "tbl":
                table = ancestor
                break
        text = Paragraph(cast("CT_P", p), parent).text
        if table is None or row is None or cell is None:
            paragraph_number += 1
            label = paragraph_label or f"paragraph {paragraph_number}"
            builder.add(text, label)
            continue
        key = tree.getpath(table)
        number = table_numbers.setdefault(key, len(table_numbers) + 1)
        row_number = table.findall(_W + "tr").index(row) + 1
        cell_number = row.findall(_W + "tc").index(cell) + 1
        builder.add(
            text, f"{table_prefix}table {number}, row {row_number}, cell {cell_number}"
        )


def _header_footer_parts(
    document: DocumentObject, section_number: int, section: Section
):
    """(label, part) for every header and footer this section defines."""
    even = document.settings.odd_and_even_pages_header_footer
    first = section.different_first_page_header_footer
    candidates = [
        (f"header {section_number}", section.header),
        (f"footer {section_number}", section.footer),
    ]
    if first:
        candidates.append(
            (f"first-page header {section_number}", section.first_page_header)
        )
        candidates.append(
            (f"first-page footer {section_number}", section.first_page_footer)
        )
    if even:
        candidates.append(
            (f"even-page header {section_number}", section.even_page_header)
        )
        candidates.append(
            (f"even-page footer {section_number}", section.even_page_footer)
        )
    # A part linked to the previous section repeats that section's content.
    return [
        (label, part) for label, part in candidates if not part.is_linked_to_previous
    ]


def _extract_docx(data: bytes) -> ExtractedText:
    builder = _Builder()
    document = docx.Document(io.BytesIO(data))
    body = cast("CT_Document", document.element).body
    _walk_wordml(body, document, builder, paragraph_label=None, table_prefix="")
    for section_number, section in enumerate(document.sections, start=1):
        for label, part in _header_footer_parts(document, section_number, section):
            _walk_wordml(
                part.part.element,
                part,
                builder,
                paragraph_label=label,
                table_prefix=f"{label}, ",
            )
    return builder.build()


def _extract_xlsx(data: bytes) -> ExtractedText:
    builder = _Builder()
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    builder.add(
                        str(cell.value), f"sheet {sheet.title}, cell {cell.coordinate}"
                    )
    finally:
        workbook.close()
    return builder.build()


def extract_text(source: str | os.PathLike[str] | bytes, kind: Kind) -> ExtractedText:
    """Extract plain text and a location map from a .docx or .xlsx file."""
    if kind not in ("docx", "xlsx"):
        raise ValueError(f"unsupported kind {kind!r}; expected 'docx' or 'xlsx'")
    data = _read_source(source)
    _check_zip(data)
    try:
        return _extract_docx(data) if kind == "docx" else _extract_xlsx(data)
    except Exception as exc:
        raise ValueError(f"could not parse {kind} document: {exc}") from exc
